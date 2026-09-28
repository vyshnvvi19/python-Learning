import openpyxl


REQUIRED_COLUMNS = [
    "Order_ID",
    "Product",
    "Quantity",
    "Price",
    "City"
]


def validate_data(sheet):
    errors = []

    headers = [cell.value for cell in sheet[1]]

    for column in REQUIRED_COLUMNS:
        if column not in headers:
            errors.append(f"Missing required column: {column}")

    if errors:
        return errors

    column_index = {
        column: headers.index(column) + 1
        for column in REQUIRED_COLUMNS
    }

    order_ids = []

    for row in range(2, sheet.max_row + 1):
        order_id = sheet.cell(row, column_index["Order_ID"]).value
        product = sheet.cell(row, column_index["Product"]).value
        quantity = sheet.cell(row, column_index["Quantity"]).value
        price = sheet.cell(row, column_index["Price"]).value

        if not order_id:
            errors.append(f"Row {row}: Order_ID is required")

        if not product:
            errors.append(f"Row {row}: Product is required")

        if not isinstance(quantity, (int, float)) or quantity <= 0:
            errors.append(
                f"Row {row}: Quantity must be greater than 0"
            )

        if not isinstance(price, (int, float)) or price <= 0:
            errors.append(
                f"Row {row}: Price must be greater than 0"
            )

        if order_id in order_ids:
            errors.append(
                f"Row {row}: Duplicate Order_ID '{order_id}'"
            )
        else:
            order_ids.append(order_id)

    return errors


def create_report(sheet, errors):
    workbook = openpyxl.Workbook()

    report = workbook.active
    report.title = "Sales Report"

    report.append([
        "Order_ID",
        "Product",
        "Quantity",
        "Price",
        "City",
        "Total_Sales"
    ])

    headers = [cell.value for cell in sheet[1]]

    column_index = {
        column: headers.index(column) + 1
        for column in REQUIRED_COLUMNS
    }

    city_totals = {}

    for row in range(2, sheet.max_row + 1):
        order_id = sheet.cell(row, column_index["Order_ID"]).value
        product = sheet.cell(row, column_index["Product"]).value
        quantity = sheet.cell(row, column_index["Quantity"]).value
        price = sheet.cell(row, column_index["Price"]).value
        city = sheet.cell(row, column_index["City"]).value

        total_sales = quantity * price

        report.append([
            order_id,
            product,
            quantity,
            price,
            city,
            total_sales
        ])

        city_totals[city] = city_totals.get(city, 0) + total_sales

    summary = workbook.create_sheet("Summary")

    summary.append(["Sales Summary"])
    summary.append([])
    summary.append(["Total Orders", sheet.max_row - 1])

    total_sales = sum(city_totals.values())

    summary.append(["Total Sales", total_sales])

    summary.append([])
    summary.append(["City", "Total Sales"])

    for city, total in city_totals.items():
        summary.append([city, total])

    error_summary = workbook.create_sheet("Error Summary")

    error_summary.append(["Error Number", "Diagnostic"])

    if errors:
        for number, error in enumerate(errors, start=1):
            error_summary.append([number, error])
    else:
        error_summary.append([
            "-",
            "No validation errors found"
        ])

    workbook.save("output/sales_report.xlsx")


def main():
    workbook = openpyxl.load_workbook(
        "input/valid_sales_input.xlsx"
    )

    sheet = workbook.active

    errors = validate_data(sheet)

    if errors:
        print("Validation failed.")

        for error in errors:
            print("-", error)

        create_report(sheet, errors)

        print("\nError summary created.")
        return

    create_report(sheet, errors)

    print("Validation passed.")
    print("Sales report created successfully.")


if __name__ == "__main__":
    main()