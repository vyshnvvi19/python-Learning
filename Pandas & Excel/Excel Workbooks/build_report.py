import openpyxl
from openpyxl.styles import Font, Alignment

workbook = openpyxl.load_workbook("sales_data.xlsx")

sales = workbook["Sales"]
products = workbook["Products"]

# Create or replace Summary sheet
if "Summary" in workbook.sheetnames:
    del workbook["Summary"]

summary = workbook.create_sheet("Summary")

# Add calculated sales column
sales["F1"] = "Total Sales"

for row in range(2, sales.max_row + 1):
    sales[f"F{row}"] = f"=C{row}*D{row}"

# Summary title
summary["A1"] = "Sales Report Summary"
summary["A1"].font = Font(bold=True, size=14)

# Basic summary
summary["A3"] = "Total Orders"
summary["B3"] = f"=COUNTA(Sales!A2:A{sales.max_row})"

summary["A4"] = "Total Sales"
summary["B4"] = f"=SUM(Sales!F2:F{sales.max_row})"

summary["A5"] = "Average Sale"
summary["B5"] = f"=AVERAGE(Sales!F2:F{sales.max_row})"

# Sales by city
summary["A7"] = "Sales by City"
summary["A8"] = "City"
summary["B8"] = "Total Sales"

cities = {}

for row in range(2, sales.max_row + 1):
    city = sales[f"E{row}"].value
    quantity = sales[f"C{row}"].value
    price = sales[f"D{row}"].value
    total = quantity * price

    cities[city] = cities.get(city, 0) + total

row_number = 9

for city, total in cities.items():
    summary[f"A{row_number}"] = city
    summary[f"B{row_number}"] = total
    row_number += 1

# Formatting
for cell in summary[1]:
    cell.alignment = Alignment(horizontal="center")

for cell in summary[8]:
    cell.font = Font(bold=True)

summary.column_dimensions["A"].width = 25
summary.column_dimensions["B"].width = 18

for row in range(2, sales.max_row + 1):
    sales[f"D{row}"].number_format = '#,##0'
    sales[f"F{row}"].number_format = '#,##0'

summary["B4"].number_format = '#,##0'
summary["B5"].number_format = '#,##0.00'

for row in range(9, row_number):
    summary[f"B{row}"].number_format = '#,##0'

# Format Sales headers
for cell in sales[1]:
    cell.font = Font(bold=True)

# Save final report
workbook.save("sales_report.xlsx")

print("Excel report created successfully.")