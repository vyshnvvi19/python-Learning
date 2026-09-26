import openpyxl


REQUIRED_COLUMNS = ["Student_ID", "Name", "Age", "Marks"]


def validate_excel(filename):
    workbook = openpyxl.load_workbook(filename)
    sheet = workbook.active

    errors = []

    headers = [cell.value for cell in sheet[1]]

    # Schema validation
    for column in REQUIRED_COLUMNS:
        if column not in headers:
            errors.append(f"Missing required column: {column}")

    if errors:
        return errors

    column_index = {
        column: headers.index(column) + 1
        for column in REQUIRED_COLUMNS
    }

    student_ids = []

    for row in range(2, sheet.max_row + 1):
        student_id = sheet.cell(row, column_index["Student_ID"]).value
        name = sheet.cell(row, column_index["Name"]).value
        age = sheet.cell(row, column_index["Age"]).value
        marks = sheet.cell(row, column_index["Marks"]).value

        if not student_id:
            errors.append(f"Row {row}: Student_ID is required")

        if not name:
            errors.append(f"Row {row}: Name is required")

        if not isinstance(age, (int, float)) or not 1 <= age <= 100:
            errors.append(f"Row {row}: Age must be between 1 and 100")

        if not isinstance(marks, (int, float)) or not 0 <= marks <= 100:
            errors.append(f"Row {row}: Marks must be between 0 and 100")

        if student_id in student_ids:
            errors.append(
                f"Row {row}: Duplicate Student_ID '{student_id}'"
            )
        else:
            student_ids.append(student_id)

    return errors


def create_error_report(errors):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Validation Errors"

    sheet.append(["Error Number", "Diagnostic"])

    for number, error in enumerate(errors, start=1):
        sheet.append([number, error])

    workbook.save("validation_errors.xlsx")


def print_validation_report(errors):
    print("VALIDATION REPORT")
    print("=================")

    if errors:
        print("Validation failed.")
        print("\nErrors:")

        for error in errors:
            print("-", error)

        create_error_report(errors)
        print("\nError report saved as validation_errors.xlsx")

    else:
        print("Validation passed. No errors found.")


if __name__ == "__main__":
    errors = validate_excel("students.xlsx")
    print_validation_report(errors)